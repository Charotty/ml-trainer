#!/usr/bin/env python3
"""Compare anatomy-extractor output with «Смещение - конечное -13».

The Excel columns below are the ones the extractor can now fill. Lateral-scan
coordinates stay out of this comparison: a supine extraction must not be scored
against the "на боку" block.

References
  TotalSegmentator classes:
    https://github.com/wasserth/TotalSegmentator/blob/master/totalsegmentator/map_to_binary.py
  MAP score (posterior fat + stranding):
    https://pubmed.ncbi.nlm.nih.gov/25192968/

Usage
  py -3 scripts/validation/validate_extractor_anatomy.py
  py -3 scripts/validation/validate_extractor_anatomy.py --extracted results/na_spine_full.csv

Without ``--extracted`` the report still lists every compared column and how
often it is filled in the workbook (status ``no_extraction``).
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path
from typing import Dict, List, Optional

import numpy as np
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.features.ct_anatomy.calibration import suggest_calibration  # noqa: E402

# 0-based columns on the «Смещение» sheet. Shared clinical fields:
SHARED_COLUMNS = {
    "age": 3,
    "body_type": 4,
    "bmi": 8,
    "abd_wall_thickness_mm": 11,
    "lumbar_lordosis_deg": 12,
    "s1_plate_tilt_deg": 13,
    "disc_l3l4_anterior_height_mm": 14,
    "disc_l3l4_posterior_height_mm": 15,
    "disc_l4l5_anterior_height_mm": 16,
    "disc_l4l5_posterior_height_mm": 17,
    "abd_depth_l3l4_mm": 18,
    "abd_width_l3l4_mm": 19,
}

# Right kidney starts at column 21. Left kidney is the same block + 42.
SIDE_FIELDS = {
    "perirenal_dorsal_mm": 21,
    "perirenal_ventral_mm": 22,
    "perirenal_lateral_mm": 23,
    "perirenal_hu": 24,
    "medial_to_spine_mm": 25,
    "map_score": 26,
    "psoas_thickness_mm": 27,
    "psoas_area_cm2": 28,
    "upper_pole_to_rib11_mm": 29,
    "lower_pole_to_iliac_crest_mm": 30,
    "upper_pole_to_diaphragm_mm": 31,
    "rotation_a_deg": 41,
    "rotation_b_deg": 42,
    "rotation_c_deg": 43,
    "pedicle_length_mm": 44,
    "pedicle_origin_angle_deg": 45,
    "pedicle_drop_angle_deg": 46,
}
LEFT_OFFSET = 42
DATA_START_ROW = 5  # 0-based, same as xlsx_displacement_parser

BODY_TYPE_MAP = {
    "норма": 0.0,
    "нормостеническое": 0.0,
    "астеническое": 1.0,
    "гиперстеническое": 2.0,
    "гипер": 2.0,
}

DEFAULT_XLSX = REPO_ROOT / "Смещение - конечное -13 .xlsx"
DEFAULT_REPORT = REPO_ROOT / "results" / "extractor_validation.csv"
DEFAULT_CALIBRATION = REPO_ROOT / "results" / "anatomy_calibration_suggested.json"


def surname_key(value: object) -> str:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return ""
    text = str(value).strip().lower().replace("ё", "е")
    token = re.split(r"[\s.]+", text)[0]
    return re.sub(r"[^a-zа-я0-9]", "", token)


def _parse_number(value: object) -> float:
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return np.nan
    if isinstance(value, (int, float, np.integer, np.floating)):
        return float(value)
    text = str(value).strip().lower().replace(",", ".").replace("\u00a0", "")
    if text in BODY_TYPE_MAP:
        return BODY_TYPE_MAP[text]
    if text in {"", "-", "nan", "none"}:
        return np.nan
    try:
        return float(text)
    except ValueError:
        return np.nan


def _cell(row: tuple, col: int) -> object:
    if col >= len(row):
        return None
    return row[col]


def read_manual_table(path: Path) -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    records: List[dict] = []
    for row in rows[DATA_START_ROW:]:
        name = surname_key(_cell(row, 1))
        if not name:
            continue
        record: Dict[str, object] = {"name_key": name}
        for feature, col in SHARED_COLUMNS.items():
            record[feature] = _parse_number(_cell(row, col))
        for side, offset in (("right", 0), ("left", LEFT_OFFSET)):
            for suffix, col in SIDE_FIELDS.items():
                record[f"kidney_{side}_{suffix}"] = _parse_number(_cell(row, col + offset))
        records.append(record)
    return pd.DataFrame(records)


def _icc_2_1(x: np.ndarray, y: np.ndarray) -> float:
    """Two-way random, absolute agreement, single measures (Shrout & Fleiss)."""
    data = np.column_stack([x, y])
    n, k = data.shape
    if n < 3:
        return float("nan")
    mean_row = data.mean(axis=1, keepdims=True)
    mean_col = data.mean(axis=0, keepdims=True)
    grand = float(data.mean())
    msr = k * float(np.var(mean_row, ddof=1))
    msc = n * float(np.var(mean_col, ddof=1))
    residual = data - mean_row - mean_col + grand
    mse = float((residual ** 2).sum() / ((n - 1) * (k - 1)))
    denom = msr + (k - 1) * mse + k * (msc - mse) / n
    if denom == 0:
        return float("nan")
    return (msr - mse) / denom


def agreement_row(feature: str, manual: np.ndarray, auto: np.ndarray) -> dict:
    """MAE, Bland–Altman bias and ICC for one paired feature."""
    ok = np.isfinite(manual) & np.isfinite(auto)
    n = int(ok.sum())
    row = {
        "feature": feature,
        "n_paired": n,
        "mae": np.nan,
        "bias": np.nan,
        "loa_low": np.nan,
        "loa_high": np.nan,
        "icc": np.nan,
        "slope": np.nan,
        "intercept": np.nan,
        "status": "ok" if n else "no_pairs",
    }
    if n == 0:
        return row
    diff = auto[ok] - manual[ok]
    bias = float(diff.mean())
    sd = float(diff.std(ddof=1)) if n > 1 else 0.0
    row["mae"] = float(np.mean(np.abs(diff)))
    row["bias"] = bias
    row["loa_low"] = bias - 1.96 * sd
    row["loa_high"] = bias + 1.96 * sd
    row["icc"] = _icc_2_1(manual[ok], auto[ok])
    suggested = suggest_calibration(manual[ok], auto[ok])
    if suggested and abs(bias) >= 2.0:
        row["slope"] = suggested["slope"]
        row["intercept"] = suggested["intercept"]
        row["status"] = "calibrate"
    return row


def compare_tables(manual: pd.DataFrame, extracted: Optional[pd.DataFrame]) -> pd.DataFrame:
    features = [c for c in manual.columns if c != "name_key"]
    if extracted is None or extracted.empty:
        rows = []
        for feature in features:
            filled = int(pd.to_numeric(manual[feature], errors="coerce").notna().sum())
            rows.append(
                {
                    "feature": feature,
                    "n_manual": filled,
                    "n_paired": 0,
                    "mae": np.nan,
                    "bias": np.nan,
                    "loa_low": np.nan,
                    "loa_high": np.nan,
                    "icc": np.nan,
                    "slope": np.nan,
                    "intercept": np.nan,
                    "status": "no_extraction",
                }
            )
        return pd.DataFrame(rows)

    frame = extracted.copy()
    name_col = next((c for c in ("patient_name", "full_name", "fio", "case_id") if c in frame.columns), None)
    if name_col is None:
        raise ValueError("extracted table has no patient_name, full_name, fio or case_id column")
    frame["name_key"] = frame[name_col].map(surname_key)
    merged = manual.merge(frame, on="name_key", how="inner", suffixes=("_manual", "_auto"))
    rows = []
    for feature in features:
        manual_col = f"{feature}_manual" if f"{feature}_manual" in merged.columns else feature
        auto_col = f"{feature}_auto" if f"{feature}_auto" in merged.columns else None
        if auto_col is None or auto_col not in merged.columns:
            filled = int(pd.to_numeric(manual[feature], errors="coerce").notna().sum())
            rows.append({**agreement_row(feature, np.array([]), np.array([])), "n_manual": filled, "status": "missing_in_extraction"})
            continue
        manual_vals = pd.to_numeric(merged[manual_col], errors="coerce").to_numpy(dtype=float)
        auto_vals = pd.to_numeric(merged[auto_col], errors="coerce").to_numpy(dtype=float)
        stats = agreement_row(feature, manual_vals, auto_vals)
        stats["n_manual"] = int(pd.to_numeric(manual[feature], errors="coerce").notna().sum())
        rows.append(stats)
    return pd.DataFrame(rows)


def _write_calibration(report: pd.DataFrame, path: Path) -> None:
    suggested = {}
    if "status" not in report.columns:
        return
    for _, row in report[report["status"] == "calibrate"].iterrows():
        if pd.notna(row["slope"]) and pd.notna(row["intercept"]):
            suggested[row["feature"]] = {"slope": float(row["slope"]), "intercept": float(row["intercept"])}
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(suggested, ensure_ascii=False, indent=2), encoding="utf-8")


def main() -> int:
    parser = argparse.ArgumentParser(description="Validate CT anatomy features against the displacement workbook")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--extracted", type=Path, default=None, help="CSV produced with --anatomy-profile full")
    parser.add_argument("--output", type=Path, default=DEFAULT_REPORT)
    parser.add_argument("--calibration-out", type=Path, default=DEFAULT_CALIBRATION)
    args = parser.parse_args()

    if not args.xlsx.exists():
        print(f"xlsx not found: {args.xlsx}")
        return 2
    manual = read_manual_table(args.xlsx)
    extracted = pd.read_csv(args.extracted) if args.extracted and args.extracted.exists() else None
    report = compare_tables(manual, extracted)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    report.to_csv(args.output, index=False)
    _write_calibration(report, args.calibration_out)
    paired = int((report["status"] == "ok").sum()) if "status" in report.columns else 0
    print(f"patients in workbook: {len(manual)}")
    print(f"features: {len(report)}  paired-ok: {paired}")
    print(f"report -> {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
