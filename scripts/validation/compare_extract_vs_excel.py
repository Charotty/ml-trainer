#!/usr/bin/env python3
"""Join CT extracts with «Смещение - конечное -13» and write the differences.

На Боку is the lateral cohort that shares surnames with the workbook.
На спине is a different list; overlap is reported, not forced.

Two coordinate readings already used in this repo are both written down:
  raw     — kidney point in the CSV (patient mm)
  vs_spine — that point minus spine_center, the fill used by normalize_dataframe
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

import numpy as np
import pandas as pd

REPO_ROOT = Path(__file__).resolve().parents[2]
if str(REPO_ROOT) not in sys.path:
    sys.path.insert(0, str(REPO_ROOT))

from src.data.holdout import normalize_surname_key  # noqa: E402
from src.data.xlsx_displacement_parser import (  # noqa: E402
    _COL,
    _cell,
    _is_header_like_row,
    _parse_numeric,
    parse_xlsx_raw_table,
)

DEFAULT_XLSX = Path("/mnt/f/Смещение - конечное -13 .xlsx")
SIDES = ("left", "right")
AXES = ("x", "y", "z")
POINTS = ("upper", "middle", "lower", "center")


def _num(frame: pd.DataFrame, col: str) -> pd.Series:
    if col not in frame.columns:
        return pd.Series(np.nan, index=frame.index, dtype=float)
    return pd.to_numeric(frame[col], errors="coerce")


def _excel_volumes(path: Path) -> pd.DataFrame:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    rows = list(wb.active.iter_rows(values_only=True))
    wb.close()
    records = []
    for sheet_idx, row in enumerate(rows):
        if sheet_idx < 5 or _is_header_like_row(row):
            continue
        surname = normalize_surname_key(_cell(row, _COL["fio"]))
        if not surname:
            continue
        records.append(
            {
                "surname": surname,
                "excel_left_volume_cm3": _parse_numeric(_cell(row, _COL["left_kidney_volume_cm3"])),
                "excel_right_volume_cm3": _parse_numeric(_cell(row, _COL["right_kidney_volume_cm3"])),
            }
        )
    volumes = pd.DataFrame(records)
    if volumes.empty:
        return volumes
    return volumes.drop_duplicates("surname", keep="first")


def load_excel(path: Path) -> pd.DataFrame:
    raw = parse_xlsx_raw_table(path, exclude_surnames=[])
    frame = raw.copy()
    frame["surname"] = frame["fio"].map(normalize_surname_key)
    frame = frame[frame["surname"].astype(str).str.len() > 0]
    counts = frame["surname"].value_counts()
    frame["excel_surname_n"] = frame["surname"].map(counts)
    volumes = _excel_volumes(path)
    if not volumes.empty:
        frame = frame.merge(volumes, on="surname", how="left")
    keep = [
        "surname",
        "fio",
        "excel_surname_n",
        "excel_row",
        "abd_width_l3l4_mm",
        "abd_depth_l3l4_mm",
        "excel_left_volume_cm3",
        "excel_right_volume_cm3",
    ]
    for side in SIDES:
        for point in ("upper", "middle", "lower"):
            for axis in AXES:
                keep.append(f"{side}_supine_{point}_{axis}")
                keep.append(f"{side}_lateral_{point}_{axis}")
                keep.append(f"{side}_delta_{point}_{axis}")
    keep = [c for c in keep if c in frame.columns]
    return frame[keep].rename(columns={"fio": "excel_fio", "excel_row": "excel_sheet_row"})


def load_extract(path: Path, position: str) -> pd.DataFrame:
    frame = pd.read_csv(path)
    frame["surname"] = frame["case_id"].map(normalize_surname_key)
    frame = frame[frame["surname"].astype(str).str.len() > 0].copy()
    frame = frame.reset_index(drop=True)
    counts = frame["surname"].value_counts()
    frame[f"{position}_surname_n"] = frame["surname"].map(counts)
    out = pd.DataFrame({"surname": frame["surname"], f"{position}_surname_n": frame[f"{position}_surname_n"]})
    out[f"{position}_case_id"] = frame["case_id"].values
    for src, dst in (
        ("status", f"{position}_status"),
        ("totalsegmentator_status", f"{position}_ts_status"),
        ("series_slices", f"{position}_series_slices"),
        ("body_width_mm", f"{position}_body_width_mm"),
        ("body_depth_mm", f"{position}_body_depth_mm"),
    ):
        if src in frame.columns:
            out[dst] = pd.to_numeric(frame[src], errors="coerce") if src != "status" and src != "totalsegmentator_status" else frame[src].values
    for side in SIDES:
        vol = f"kidney_{side}_volume_cm3"
        if vol in frame.columns:
            out[f"{position}_{side}_volume_cm3"] = pd.to_numeric(frame[vol], errors="coerce").values
        for axis in AXES:
            spine = pd.to_numeric(frame.get(f"spine_center_{axis}"), errors="coerce")
            for point in POINTS:
                src = f"kidney_{side}_{point}_{axis}"
                if src not in frame.columns:
                    continue
                raw = pd.to_numeric(frame[src], errors="coerce")
                out[f"{position}_{side}_{point}_{axis}"] = raw.values
                out[f"{position}_{side}_{point}_{axis}_vs_spine"] = (raw - spine).values
    return out


def _error(ct: pd.Series, excel: pd.Series) -> pd.Series:
    both = ct.notna() & excel.notna()
    out = pd.Series(np.nan, index=ct.index, dtype=float)
    out.loc[both] = ct.loc[both] - excel.loc[both]
    return out


def build_patient_table(excel: pd.DataFrame, spine: pd.DataFrame, boku: pd.DataFrame) -> pd.DataFrame:
    excel_u = excel[excel["excel_surname_n"] == 1]
    spine_u = spine[spine["spine_surname_n"] == 1]
    boku_u = boku[boku["boku_surname_n"] == 1]
    merged = excel_u.merge(spine_u, on="surname", how="outer")
    merged = merged.merge(boku_u, on="surname", how="outer")

    for side in SIDES:
        excel_vol = _num(merged, f"excel_{side}_volume_cm3")
        for position, excel_pose in (("boku", "lateral"), ("spine", "supine")):
            merged[f"{position}_minus_excel_{side}_volume"] = _error(
                _num(merged, f"{position}_{side}_volume_cm3"), excel_vol
            )
            for point in POINTS:
                excel_point = "middle" if point == "center" else point
                for axis in AXES:
                    excel_col = f"{side}_{excel_pose}_{excel_point}_{axis}"
                    raw_col = f"{position}_{side}_{point}_{axis}"
                    rel_col = f"{position}_{side}_{point}_{axis}_vs_spine"
                    merged[f"{position}_minus_excel_{side}_{point}_{axis}"] = _error(
                        _num(merged, raw_col), _num(merged, excel_col)
                    )
                    merged[f"{position}_vs_spine_minus_excel_{side}_{point}_{axis}"] = _error(
                        _num(merged, rel_col), _num(merged, excel_col)
                    )
        for feature, col in (
            ("abd_width_l3l4_mm", "body_width_mm"),
            ("abd_depth_l3l4_mm", "body_depth_mm"),
        ):
            for position in ("boku", "spine"):
                merged[f"{position}_minus_excel_{feature}"] = _error(
                    _num(merged, f"{position}_{col}"), _num(merged, feature)
                )

    def status_row(row: pd.Series) -> str:
        has_excel = pd.notna(row.get("excel_fio")) and str(row.get("excel_fio")).strip() not in {"", "nan"}
        has_spine = pd.notna(row.get("spine_case_id")) and str(row.get("spine_case_id")).strip() not in {"", "nan"}
        has_boku = pd.notna(row.get("boku_case_id")) and str(row.get("boku_case_id")).strip() not in {"", "nan"}
        if has_excel and has_spine and has_boku:
            return "excel_and_both_ct"
        if has_excel and has_boku:
            return "excel_and_boku"
        if has_excel and has_spine:
            return "excel_and_spine"
        if has_excel:
            return "excel_only"
        if has_spine and has_boku:
            return "both_ct_only"
        return "ct_only"

    merged["match_status"] = merged.apply(status_row, axis=1)
    return merged.sort_values(["match_status", "surname"]).reset_index(drop=True)


def duplicate_rows(excel: pd.DataFrame, spine: pd.DataFrame, boku: pd.DataFrame) -> pd.DataFrame:
    pieces = []
    for label, frame, name_col, n_col in (
        ("excel", excel, "excel_fio", "excel_surname_n"),
        ("spine", spine, "spine_case_id", "spine_surname_n"),
        ("boku", boku, "boku_case_id", "boku_surname_n"),
    ):
        dup = frame[frame[n_col] > 1]
        if dup.empty:
            continue
        piece = dup[["surname", name_col]].copy()
        piece["source"] = label
        piece = piece.rename(columns={name_col: "name"})
        pieces.append(piece)
    if not pieces:
        return pd.DataFrame(columns=["surname", "source", "name"])
    return pd.concat(pieces, ignore_index=True)


def _mae_row(patients: pd.DataFrame, mask: pd.Series, metric: str, err_col: str) -> dict:
    err = _num(patients.loc[mask], err_col)
    ok = err.dropna()
    return {
        "metric": metric,
        "n": int(ok.shape[0]),
        "mae": float(ok.abs().mean()) if len(ok) else np.nan,
        "bias_ct_minus_excel": float(ok.mean()) if len(ok) else np.nan,
        "median_abs": float(ok.abs().median()) if len(ok) else np.nan,
    }


def summary_table(patients: pd.DataFrame) -> pd.DataFrame:
    rows = []
    groups = {
        "boku_lateral": patients["match_status"].isin(["excel_and_boku", "excel_and_both_ct"]),
        "spine_supine": patients["match_status"].isin(["excel_and_spine", "excel_and_both_ct"]),
    }
    for group, mask in groups.items():
        position = "boku" if group.startswith("boku") else "spine"
        for side in SIDES:
            rows.append(
                _mae_row(
                    patients,
                    mask,
                    f"{group}/{side}_volume_cm3",
                    f"{position}_minus_excel_{side}_volume",
                )
            )
            for point in ("middle", "center"):
                for axis in AXES:
                    rows.append(
                        _mae_row(
                            patients,
                            mask,
                            f"{group}/{side}_{point}_{axis}_raw",
                            f"{position}_minus_excel_{side}_{point}_{axis}",
                        )
                    )
                    rows.append(
                        _mae_row(
                            patients,
                            mask,
                            f"{group}/{side}_{point}_{axis}_vs_spine",
                            f"{position}_vs_spine_minus_excel_{side}_{point}_{axis}",
                        )
                    )
        for feature in ("abd_width_l3l4_mm", "abd_depth_l3l4_mm"):
            rows.append(
                _mae_row(
                    patients,
                    mask,
                    f"{group}/{feature}",
                    f"{position}_minus_excel_{feature}",
                )
            )
    return pd.DataFrame(rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare CT extract CSVs with the displacement workbook")
    parser.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    parser.add_argument("--spine", type=Path, default=REPO_ROOT / "results" / "na_spine_full.csv")
    parser.add_argument("--boku", type=Path, default=REPO_ROOT / "results" / "na_boku_full.csv")
    parser.add_argument("--patients-out", type=Path, default=REPO_ROOT / "results" / "ct_vs_excel_patients.csv")
    parser.add_argument("--summary-out", type=Path, default=REPO_ROOT / "results" / "ct_vs_excel_summary.csv")
    parser.add_argument("--duplicates-out", type=Path, default=REPO_ROOT / "results" / "ct_vs_excel_duplicate_surnames.csv")
    args = parser.parse_args()

    if not args.xlsx.is_file():
        print(f"xlsx not found: {args.xlsx}")
        return 2

    excel = load_excel(args.xlsx)
    spine = load_extract(args.spine, "spine")
    boku = load_extract(args.boku, "boku")
    patients = build_patient_table(excel, spine, boku)
    summary = summary_table(patients)
    duplicates = duplicate_rows(excel, spine, boku)

    args.patients_out.parent.mkdir(parents=True, exist_ok=True)
    patients.to_csv(args.patients_out, index=False)
    summary.to_csv(args.summary_out, index=False)
    duplicates.to_csv(args.duplicates_out, index=False)

    print(f"excel rows: {len(excel)}  spine: {len(spine)}  boku: {len(boku)}")
    print(patients["match_status"].value_counts().to_string())
    show = summary[summary["metric"].str.contains("middle_|volume_cm3|abd_")]
    print(show.to_string(index=False))
    print(f"patients   -> {args.patients_out}")
    print(f"summary    -> {args.summary_out}")
    print(f"duplicates -> {args.duplicates_out}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
